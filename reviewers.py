import openreview
import pandas as pd
import numpy as np
from tqdm import tqdm
from collections import defaultdict
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity
import itertools
import re
import pprint

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill
red_fill = PatternFill( fill_type="solid", start_color="FFC7CE", end_color="FFC7CE" )
green_fill = PatternFill(fill_type="solid", start_color="C6EFCE", end_color="C6EFCE" )
yellow_fill = PatternFill(fill_type="solid", start_color="FFEB9C", end_color="FFEB9C" )
from openpyxl.styles import Alignment
center = Alignment( horizontal="center", vertical="center" )
center_wrap = Alignment( horizontal="center", vertical="center", wrap_text = True )

import config_ivo as c
from collections import defaultdict
import time

from Lists.IWAI2026_Exclude import EXCLUDE_REVIEWERS
from Lists.IWAI2026_Exclude import EXCLUDE_PAPERS

# SUBMISSION TYPES
streams = ['1-comp','2-cogn','3-appl']
types=['1-paper','2-abstr']
TECHNICAL_STREAM = streams[0] # Computational
only_papers=True

BASEURL = 'https://api2.openreview.net'
YEAR = '2026'
VENUE = f'IWAI-{YEAR}'
venue_id = f'IWAI/{YEAR}/Workshop'
SUBMISSION_INVITATION = f'{venue_id}/-/Submission'
ASSIGNMENT_INVITATION = f'{venue_id}/Reviewers/-/Assignment'
COI_INVITATION = f"{venue_id}/Reviewers/-/Conflict"
REVIEWER_GROUP = f'{venue_id}/Reviewers'
REVIEW_INVITATION = f'{venue_id}/Submission*/-/Official_Review'
MSG_INVITATION=f'{venue_id}/-/Edit'


client=openreview.api.OpenReviewClient(baseurl=BASEURL,username=c.usr,password=c.pas)
submissions = client.get_all_notes(invitation=SUBMISSION_INVITATION)
print(f"Downloaded {len(submissions)} submissions.")

assigned_reviews = defaultdict(set)
completed_reviews = defaultdict(set)
missing_reviews = defaultdict(set)

#reviewer2anonym, anonym2reviewer = {}, {}
from Lists.iwai2026_anonym2reviewers import anonym2reviewer

def extract_field(content, key):
    """  OpenReview V2 stores values as:  content[key]["value"] or content[key]  """
    if key not in content:  return None
    value = content[key]
    if isinstance(value, dict) and "value" in value:
        return value["value"]
    return value

def reviewers_and_authors():
    # paper_reviewers = client.get_group(f"{VENUE}/Submission31/Reviewers")
    # print(paper_reviewers.members)

    ass1,ass2 = [],[]
    maxr1, maxr2, maxa, maxr = 0,0,0,0
    reviewer_to_subs = defaultdict(list)
    pap={}
    for s in submissions:
        paper_id=s.id
        number = getattr(s, "number", 0)
        #if not number == 31: continue
        typ=s.content.get("type")
        typ = int(typ["value"][0]) if typ else 1 # Fallback for previous IWAI editions without Type.
        strm=s.content.get("stream")["value"]
        stream = streams[int(strm[0])-1]
        title= s.content['title']['value'] if isinstance(s.content['title'], dict) else s.content['title']
        authors = s.content['authorids']['value']
        maxa = max(maxa,len(authors))

        edges=client.get_edges(invitation=ASSIGNMENT_INVITATION, head=paper_id)
        reviewers = [edge.tail for edge in edges]

        for r in reviewers: reviewer_to_subs[r].append(number)
        maxr = max(maxr, len(reviewers))

        entry = {'ID':paper_id, 'number':number, 'stream':stream, 'title':title, 'reviewers':reviewers, 'authors':authors }
        if typ==1:
            ass1.append(entry)
            maxr1=max(maxr1,len(reviewers))
        else:
            ass2.append(entry)
            maxr2=max(maxr2,len(reviewers))

    ass1.sort(key=lambda x: (x['stream'], x['number']))
    ass2.sort(key=lambda x: (x['stream'], x['number']))

    # Create workbook
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Papers"
    ws2 = wb.create_sheet("Posters")
    def format_reviewer(rid):
        ass_pap = reviewer_to_subs[rid]
        return f"{rid} - {len(ass_pap)}\n[{','.join(str(p) for p in ass_pap)}]" # "" for new line
    def format_author(aid):
        #return f"{aid}"
        ass_pap = reviewer_to_subs.get(aid, [])
        if ass_pap:
            return f"{aid} - {len(ass_pap)}\n[{','.join(str(p) for p in ass_pap)}]"
        else:
            return aid
    def write_sheet(ws, assignment_list):
        header = ( ["ID", "stream", "title"]
            + [f"REV_{i+1}" for i in range(4)] + [f"AUT_{i+1}" for i in range(10)])
        ws.append(header)
        for a in assignment_list:
            revs = [format_reviewer(r) for r in a['reviewers'][:4]]
            revs += [""] * (4 - len(revs))
            auts = [format_author(auth) for auth in a['authors'][:10]]
            auts += [""] * (10 - len(auts))
            row = [a['number'], a['stream'], a['title']] + revs + auts
            ws.append(row)

    write_sheet(ws1, ass1);
    write_sheet(ws2, ass2)
    filename=f"{VENUE}-reviewers-and-authors--.xlsx"
    wb.save(filename)
    print(f"Exported reviewer assignments and authors to {filename}")

def xx___reviewers_with_pending_reviews():
    pending = {}
    for s in submissions:
        number = getattr(s, "number", 0)
        paper_id = s.id
        # paper_group = f"{venue_id}/Paper{paper_number}"
        # reviewers_group = f"{paper_group}/Reviewers"
        # review_inv = f"{paper_group}/-/Official_Review"

        # Get assigned reviewers via edges
        edges = client.get_edges(invitation=ASSIGNMENT_INVITATION, head=paper_id)
        reviewers = [edge.tail for edge in edges]

        # Extract submitted reviews for this paper only
        review_replies=[]
        if isinstance(s.details, dict) and "replies" in s.details:
            replies = s.details['replies']
            review_replies = [
                r for r in replies
                if any(inv.endswith("Official_Review") for inv in r['invitations'])
            ]

        # Resolve actual reviewer IDs from signatures (may be anonymous groups)
        submitted_reviewers = set()
        for r in review_replies:
            signature = r.get('signatures', [None])[0]
            if signature:
                try:
                    group = client.get_group(id=signature)
                    submitted_reviewers.update(group.members)
                except:
                    submitted_reviewers.add(signature)


        # Find missing reviewers
        missing = [r for r in reviewers if r not in submitted_reviewers]
        if missing:
            pending[f'Paper{number}'] = missing

    return pending

def export_reviews_():
    """ Export review summary per submission. Each row is a review with metadata. """

    Rpap,Rabs = [],[]
    for s in submissions:
        number = s.number
        title = extract_field(s.content, 'title')
        type = extract_field(s.content,'type');     type_int = int(type[0])
        stream = extract_field(s.content,'stream'); Stream = streams[int(stream[0])-1]

        reviews = client.get_notes(invitation=REVIEW_INVITATION.replace('*', str(number)))

        for idx, review in enumerate(reviews):
            rating = review.content.get('rating', None)["value"]
            confidence = review.content.get('confidence', None)["value"]
            summary = review.content.get('title', '')["value"]
            review_txt  = review.content.get('review', '')
            reviewer = review.signatures[0] if review.signatures else 'N/A'
            row={
                'sNumber': number,
                'sTitle':  title,
                'Stream':  Stream,
                'Reviewer': reviewer,
                'Review #': idx + 1,
                'Review Title': summary,
                'Rating':  rating,
                'Confidence': confidence
            }
            if type_int==1: Rpap.append(row)
            else:           Rabs.append(row)

    pd.DataFrame(Rpap).to_excel('review_export_1.xlsx', index=False)
    pd.DataFrame(Rabs).to_excel('review_export_2.xlsx', index=False)

def export_reviews(fname="IWAI-ReviewScores.xlsx"):
    wb = Workbook()
    ws1 = wb.active; ws1.title = '1.Papers'
    ws2 = wb.create_sheet('2.Abstr')
    # Header
    base_cols = ['Numb','Stream','Title','Auths']
    max_reviews = 3
    review_cols = ['SCORE']
    for i in range(1, max_reviews + 1): review_cols += [f'R{i}']
    headers = base_cols + review_cols + ["Rate (subscales) Confid", "Reviewer", "Summary"]
    ws1.append(headers); ws2.append(headers)

    # Reviews
    # max_n=135; SCORES=np.zeros((max_n,7,max_reviews))
    for s in tqdm(sorted(submissions, key=lambda x: int(x.number)),
                  total=len(submissions), desc="Export reviews"): # Elaborate, sorted by submission number
        number = s.number
        title = extract_field(s.content, 'title')
        authors = extract_field(s.content,"authors")
        type = extract_field(s.content,'type');     type_int = int(type[0])
        stream = extract_field(s.content,'stream'); Stream = streams[int(stream[0])-1]
        row = [number, Stream, title, "\n".join(authors)]
        reviews = client.get_notes(invitation=REVIEW_INVITATION.replace('*', str(number)))
        Scores, ScoreStr, ReviewerStr, ReviewStr=[],"","",""

        # Weights of the five sub-criteria and for the overall rating
        Wsc = np.array([30, 25, 20, 15, 10, 50], dtype=float)
        #Wsc = np.array([30, 20, 20, 15, 15, 50], dtype=float)
        normWsc = Wsc / sum(Wsc)                                # Normalize the weights to 0-1 scale
        MaxScores = np.array([5,5,5,5,5,6],dtype=float)     # Max-score of each subscale (soun, orig, sign, pert, pres, rate)

        for idx in range(max_reviews):
            if idx: # Add newline for 2nd+ rows
                ScoreStr += "\n"; ReviewStr += "\n"; ReviewerStr +="\n"
            if idx < len(reviews):
                review = reviews[idx]
                r_title = extract_field(review.content,'title')
                #r_review = review.content.get('review', '')
                soun = extract_field(review.content,'soundness')    # 30
                orig = extract_field(review.content,'originality')  # 25
                sign = extract_field(review.content,'significance') # 20
                pert = extract_field(review.content,'pertinence')   # 15
                pres = extract_field(review.content,'presentation') # 10    # Sub-criteria weights; 100 in total
                rate = extract_field(review.content,'rating')       # 50    # Overall rating weights 50
                conf = extract_field(review.content,'confidence')           # Confidence is used to shrink towards mid-scale
                scores = [soun,orig,sign,pert,pres,rate,conf] # Scale 1-5 for each sub-criteria and 1-6 for Overall Rating

                # Compute overall Reviewer Score weighted by importance of each subscale and shrinked towards mid-scale according to uncertainty
                scores_=np.array(scores[:-1], dtype=float)  # The scores without the confidence
                scores01=(scores_-1)/(MaxScores-1)          # Normalize all scales to [0-1]
                WSC = sum(scores01*normWsc)                 # Overall normalized weighted score on a 0-1 scale
                normConf = conf/5                           # Normalize confidence to 0-1, starting with 0.2 for the lowest confidence
                #normConf = (conf+1)/6                       # Normalize confidence to 0-1; depart from 0.33 to avoid negative bias in self-confidence
                ConfScore = 0.5 + (WSC-0.5)*normConf        # Rescale the overall weighted vote towards mid-scale (0.5) for low confident (uncertaint) reviews
                ReviewScore = np.round(ConfScore*100)       # Overall normalized reviewer vote on a 0-100 scale
                Scores.append(ReviewScore)

                reviewer = review.signatures[0] if review.signatures else 'N/A'
                reviewer_anonym_id = reviewer.split("/")[-1] # Just Reviewer_XXXX
                reviewer_id = anonym2reviewer[reviewer] # The map includes full paper specification
                ScoreStr+=f"{rate} (so{soun} or{orig} sg{sign} pe{pert} pr{pres}) [{conf}]"
                ReviewStr +=f"- {r_title}"
                ReviewerStr += f"{reviewer_id[1:-1]}"
            else:
                Scores.append(np.nan)
                #ScoreStr+=f"{idx+1}"
        mScore=np.nanmean(np.array(Scores,dtype='double'))
        row += [round(mScore,1) if not np.isnan(mScore) else np.nan] + Scores + [ScoreStr, ReviewerStr, ReviewStr]

        def cell_col(x):
            if x<50: return red_fill
            if x<60: return yellow_fill
            return green_fill

        if type_int == 1:
            ws1.append(row); ws1.cell(ws1.max_row, 5).fill = cell_col(mScore)
        else:
            ws2.append(row); ws2.cell(ws2.max_row, 5).fill = cell_col(mScore)

    def format_sheet(ws):
        ws.freeze_panes = "A2";
        ws.auto_filter.ref = ws.dimensions
        widths = {"C": 50, "D": 24, "I":28, "J": 24, "K": 120}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = 45
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = center

    format_sheet(ws1)
    format_sheet(ws2)
    wb.save(fname)
    print(f"Saved to {fname}")

def reviewer_decode():
    for s in tqdm(submissions, total=len(submissions), desc="Reviewer-decode"):
        groups = client.get_groups(prefix=f"{venue_id}/Submission{s.number}/Reviewer_")
        for g in groups:
            for member in g.members:
                #reviewer2anonym[member] = g.id
                anonym2reviewer[g.id] = member
    print(anonym2reviewer)
    pass

def review_status(export_file="IWAI-ReviewStatus.xlsx"):
    for s in tqdm(sorted(submissions, key=lambda x: int(x.number)),
                  total=len(submissions), desc="Checking Review status"): # Elaborate, sorted by submission number
        paper_id=s.id
        paper_number = s.number
        typ=int(extract_field(s.content,"type")[0])
        reviews = client.get_notes(invitation=f"{venue_id}/Submission{paper_number}/-/Official_Review")  # All reviews accessible through anonymous ID
        edges = client.get_edges(invitation=ASSIGNMENT_INVITATION, head=paper_id)   # All review assignments (accessible with reviewer ID)
        reviewers = [edge.tail for edge in edges]

        # Simplification
        review_signatures = { r.signatures[0] for r in reviews if r.signatures }
        groups = client.get_groups(prefix=f"{venue_id}/Submission{paper_number}/Reviewer_")
        anon_map = {}
        for g in groups:
            for member in g.members:  anon_map[member] = g.id

        for edge in edges:
            reviewer = edge.tail
            assigned_reviews[reviewer].add(paper_number)
            anon = anon_map.get(reviewer)
            if anon in review_signatures:   completed_reviews[reviewer].add(paper_number)
            else:                           missing_reviews[reviewer].add(paper_number)

        # reviewer_anon_groups = client.get_groups(prefix=f"{venue_id}/Submission{paper_number}/Reviewer_", member=reviewer_id)
        # reviewer_anon_id = reviewer_anon_groups[0].id if reviewer_anon_groups else None
        # has_review = any( reviewer_anon_id in r.signatures for r in reviews) if reviewer_anon_id else False

    # Export
    wb = Workbook()
    ws = wb.active
    ws.title = "Review Status"
    ws.append(["Reviewer","Completed","Missing","nCompleted","nMissing"])
    all_reviewers = sorted( set(assigned_reviews) | set(completed_reviews) | set(missing_reviews) )

    for reviewer in all_reviewers:
        completed = sorted(completed_reviews.get(reviewer, []))
        missing = sorted(missing_reviews.get(reviewer, []))
        ws.append([ reviewer,
            ", ".join(map(str, sorted(completed))),
            ", ".join(map(str, sorted(missing))),
            len(completed), len(missing) ])
        if len(missing) > 0:
            ws.cell(ws.max_row, 3).fill = red_fill  # Missing Reviews
            ws.cell(ws.max_row, 5).fill = red_fill  # # Missing

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(export_file)
    print(f"Saved to {export_file}")

def send_review_reminders(review_status_file, dry_run=True):
    """ Reads review status from IWAIxxxx-ReviewStatus.xlsx and sends emails """
    wb = load_workbook(review_status_file, data_only=True)
    ws = wb.active
    sent=0
    for reviewer, completed, missing, n_completed, n_missing in ws.iter_rows(min_row=2, values_only=True):
        n_completed = int(n_completed or 0)
        n_missing = int(n_missing or 0)
        total = n_completed + n_missing
        if not n_missing: continue
        missing_papers = [p.strip() for p in str(missing).split(",") if p.strip()]
        subject = f"{VENUE} Reminder: {n_missing} pending review{'s' if n_missing>1 else ''} due by 5 July"
        message = f"""Dear Reviewer,

Thank you for serving as a reviewer for {VENUE}. Your contribution is essential to maintaining the quality of the conference programme.
"""
        if n_completed>0:
            message+=f"\nYou have completed {n_completed} out of {total} assigned reviews.\n"

        message+=f"""
You still have {n_missing} pending review{'s' if n_missing > 1 else ''}. We remind you that the reviews are due by 5 July.

The outstanding submission{'s are' if n_missing > 1 else ' is'}:
"""
        for paper in missing_papers: message += f" • Submission #{paper}\n"
        message += f"""
We would greatly appreciate it if you could submit your review(s) in due time.

Thank you very much for your contribution.

Best regards,
The {VENUE} Program Chairs
"""

        if dry_run:
            if dry_run=="print":
                print("=" * 70); print(reviewer); print(subject); print(message)
            else:
                print(f"Dry test: TO {reviewer} SBJ {subject}")
        else:
            print(f"Submitting TO {reviewer} SBJ={subject}")
            client.post_message(subject=subject, message=message, recipients=[reviewer],signature=c.usr, invitation=MSG_INVITATION,)
            time.sleep(0.3)
            sent += 1

    if not dry_run: print(f"Sent {sent} reminder messages.")

def send_second_review_reminders(review_status_file, dry_run=True):
    """ Reads review status from IWAIxxxx-ReviewStatus.xlsx and sends emails """
    wb = load_workbook(review_status_file, data_only=True)
    ws = wb.active
    sent=0
    for reviewer, completed, missing, n_completed, n_missing in ws.iter_rows(min_row=2, values_only=True):
        n_completed = int(n_completed or 0)
        n_missing = int(n_missing or 0)
        total = n_completed + n_missing
        if not n_missing: continue
        missing_papers = []
        if missing:
            missing_papers = [p.strip() for p in str(missing).split(",") if p.strip()]
        name = reviewer[1:-1].split("_")[0]
        subject = f"URGENT: {VENUE}, {n_missing} pending review{'s' if n_missing>1 else ''} overdue"
        message = f"""Dear {name},

Thank you for serving as a reviewer for {VENUE}. 

Your timely contribution is essential for ensuring a high-quality technical programme and allowing us to complete the decision process on schedule.
"""
        if n_completed>0:
            message+=f"\nYou have completed {n_completed} out of {total} assigned reviews.\n"

        message+=f"""
However, you still have {n_missing} pending review{'s' if n_missing > 1 else ''} past the review deadline of 5 July.

The outstanding submission{'s are' if n_missing > 1 else ' is'}:\n
"""
        for paper in missing_papers: message += f" • Submission #{paper}\n"
        message += f"""
We urgently ask you to submit your outstanding review(s) within the next two days.

In the mean time, please reply to this email within four hours to confirm that you will complete the outstanding review(s).

Thank you very much for your contribution.

Best regards,
The {VENUE} Program Chairs
"""

        if dry_run:
            if dry_run=="print":
                print("=" * 70); print(reviewer); print(subject); print(message)
            else:
                print(f"Dry test: TO {reviewer} SBJ {subject}")
        else:
            print(f"Submitting TO {reviewer} SBJ={subject}")
            client.post_message(subject=subject, message=message, recipients=[reviewer],
                                signature=c.usr, replyTo="ivilinpeev.stoianov@cnr.it", invitation=MSG_INVITATION,)
            time.sleep(0.3)
            sent += 1
        pass

    if not dry_run: print(f"Sent {sent} reminder messages.")

def send_second_review_reminder_corr(review_status_file, dry_run=True):
    """ Reads review status from IWAIxxxx-ReviewStatus.xlsx and sends emails """
    wb = load_workbook(review_status_file, data_only=True)
    ws = wb.active
    sent=0
    for reviewer, completed, missing, n_completed, n_missing, will_do in ws.iter_rows(min_row=2, values_only=True):
        n_completed = int(n_completed or 0)
        n_missing = int(n_missing or 0)
        total = n_completed + n_missing
        if not n_missing: continue
        missing_papers = []
        if missing:
            missing_papers = [p.strip() for p in str(missing).split(",") if p.strip()]
        name = reviewer[1:-1].split("_")[0]
        subject = f"URGENT: {VENUE}, {n_missing} pending review{'s' if n_missing>1 else ''} overdue"
        message = f"""Dear {name},

A quick clarification regarding my previous email:

To confirm that you will complete your outstanding review(s), 
please reply as soon as possible directly to ivilinpeev.stoianov@cnr.it

Thank you again for your help.

Best regards,
The {VENUE} Program Chairs
"""
        if dry_run:
            print("=" * 70); print(reviewer); print(subject); print(message)
        else:
            print(f"Submitting TO {reviewer} SBJ={subject}")
            client.post_message(subject=subject, message=message, recipients=[reviewer],
                                signature=c.usr, replyTo="ivilinpeev.stoianov@cnr.it", invitation=MSG_INVITATION,)
            time.sleep(0.2)
            sent += 1
        pass

    if not dry_run: print(f"Sent {sent} reminder messages.")


if __name__ == '__main__':
    #reviewers_and_authors()
    #reviewer_decode() # Produces anonym2reviewer saved as python script to be used as shortcut

    reviews_status_file = f"{VENUE}-ReviewStatus_10Jul.xlsx"
    #review_status(reviews_status_file)
    # send_second_review_reminder_corr(reviews_status_file, dry_run=False)

    export_reviews(fname=f'{VENUE}-Reviews11_14Lug_1930.xlsx')       # and extract completed reviews
    pass